import collections 
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Create a presentation object
prs = Presentation()

# Helper to add a slide with title and bullet points
def add_slide(prs, title, bullet_points):
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    
    title_shape = slide.shapes.title
    title_shape.text = title
    
    body_shape = slide.shapes.placeholders[1]
    tf = body_shape.text_frame
    
    for i, point in enumerate(bullet_points):
        p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        p.text = point
        p.level = 0
        p.font.size = Pt(24)

# Slide 1 - Title Slide
slide_layout = prs.slide_layouts[0] # Title Slide
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "Home Bakers Product Marketplace System"
subtitle.text = "M.I.F Mifrah\nCourse / Institution\nMid Presentation"

# Slide 2 - Introduction
add_slide(prs, "Introduction", [
    "What is this project? A marketplace for home bakers.",
    "Real-life problem: Home bakers currently sell through social media.",
    "Why an online system is needed: Better organization, easier shopping, and growth for bakers."
])

# Slide 3 - Problem Statement
add_slide(prs, "Problem Statement", [
    "No proper platform dedicated to home bakers.",
    "Orders are handled manually through messages or calls.",
    "No proper tracking system for orders.",
    "Payment confusion and missing records.",
    "Difficulty in managing customers and deliveries."
])

# Slide 4 - Proposed Solution
add_slide(prs, "Proposed Solution", [
    "A simple online marketplace system.",
    "Bakers can upload their products.",
    "Customers can browse and order easily.",
    "Orders and payments are managed in one simple system.",
    "Delivery is handled through reliable delivery partners."
])

# Slide 5 - Main Features
add_slide(prs, "Main Features", [
    "User login and registration.",
    "Product browsing and search.",
    "Add to cart and place order.",
    "Order tracking system.",
    "Admin manages users and orders.",
    "Delivery partner system."
])

# Slide 6 - How the System Works
add_slide(prs, "How the System Works", [
    "Customer → Browse Products",
    "→ Add to Cart",
    "→ Place Order",
    "→ Payment",
    "→ Admin Assigns Delivery Partner",
    "→ Delivery",
    "→ Customer Receives Order"
])

# Slide 7 - Conclusion
add_slide(prs, "Conclusion", [
    "Project progress so far is on track.",
    "Benefits: Easier for home bakers to sell and customers to buy.",
    "Future improvements: Mobile app, online payments, and live tracking."
])

prs.save("Home_Bakers_Marketplace.pptx")
print("Presentation saved as Home_Bakers_Marketplace.pptx")
